"""Excel parsing for the course allotment workbook.

Three sheets are consumed:

``Subject wise``   authoritative grid  ->  Subject | Semester | Section | Faculty
``Allotted New``   faculty load sheet  ->  Faculty | subjects per semester
``Courses``        credit catalogue    ->  Course code | Course name | Credits

The parser is defensive on purpose: headers, merged cells, NaN values, ``-``
placeholders, stray numeric cells and duplicated rows are all tolerated.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .config import TimetableConfig
from .models import (
    Assignment,
    Course,
    FacultyProfile,
    Section,
    semester_ordinal,
)

SEC_RE = re.compile(r"^SEC\.?\s*([A-F])$", re.IGNORECASE)
ROMAN_RE = re.compile(r"^(I{1,3}|IV|VI{0,3}|IX|X)$", re.IGNORECASE)
COURSE_CODE_RE = re.compile(r"^[A-Z0-9]{7,14}$")
PLACEHOLDER_NAMES = {"-", "--", "--", "n/a", "na", "nil", "none", "tbd", "?", ""}
LAB_TOKEN = "lab"

# Whole-string aliases applied before tokenisation.
SUBJECT_ALIASES: Dict[str, str] = {
    "uhv": "universal human values",
    "rmipr": "research methodology and ipr",
    "rmipr research methodology ipr": "research methodology and ipr",
    "ds lab": "data structures laboratory",
    "ds": "data structures",
    "management and entrepreneurship for it": "management and engineering economics",
    "open elective basics of ai or ml": "open elective basics of ai ml",
    "open elective basics of ai/ml": "open elective basics of ai ml",
}

CONTRACTIONS = {
    "it's": "its",
    "its'": "its",
    "oops": "oops",
}

STOP_WORDS = {
    "and",
    "or",
    "of",
    "in",
    "for",
    "to",
    "the",
    "its",
    "it",
    "a",
    "an",
    "with",
    "on",
    "is",
    "at",
    "by",
}

ROMAN_VALUES = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
    "VI": 6,
    "VII": 7,
    "VIII": 8,
    "IX": 9,
    "X": 10,
}


# ----------------------------------------------------------------------
# small text helpers
# ----------------------------------------------------------------------
def clean_text(value: object) -> str:
    """Normalise a cell into a trimmed string; NaN/None become ''."""

    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    text = str(value).replace("\u00a0", " ").strip()
    if text.lower() in {"nan", "none"}:
        return ""
    return text


def is_blank_marker(value: object) -> bool:
    text = clean_text(value).lower()
    return text in PLACEHOLDER_NAMES


def normalise_name(name: str) -> str:
    text = clean_text(name).lower().replace("&", " and ")
    for wrong, right in CONTRACTIONS.items():
        text = re.sub(rf"\b{re.escape(wrong)}\b", right, text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def subject_key(name: str) -> frozenset:
    """Token signature of a subject, used for fuzzy course matching."""

    flat = normalise_name(name)
    flat = SUBJECT_ALIASES.get(flat, flat)
    tokens = [t for t in flat.split() if t and t not in STOP_WORDS]
    stemmed = {t[:-1] if len(t) > 3 and t.endswith("s") else t for t in tokens}
    return frozenset(stemmed)


def jaccard(left: frozenset, right: frozenset) -> float:
    if not left or not right:
        return 0.0
    return len(left & right) / len(left | right)


def parse_semester(raw: object) -> Optional[int]:
    """Extract the semester number from '1st SEM', 'III', '5' ..."""

    text = clean_text(raw)
    if not text:
        return None
    digits = re.search(r"\d+", text)
    if digits:
        return int(digits.group())
    roman = ROMAN_RE.match(text.replace(" SEM", "").replace("SEM", "").strip())
    if roman:
        return ROMAN_VALUES.get(roman.group().upper())
    return None


def is_lab_subject(name: str) -> bool:
    return LAB_TOKEN in normalise_name(name)


# ----------------------------------------------------------------------
# parse results
# ----------------------------------------------------------------------
@dataclass
class SubjectRow:
    subject: str
    semester: int
    sections: Dict[str, str] = field(default_factory=dict)


@dataclass
class ParseResult:
    courses: List[Course] = field(default_factory=list)
    subject_rows: List[SubjectRow] = field(default_factory=list)
    assignments: List[Assignment] = field(default_factory=list)
    profiles: Dict[str, FacultyProfile] = field(default_factory=dict)
    hours_report: Dict[str, Tuple[float, str]] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    unmapped_subjects: List[Tuple[str, str]] = field(default_factory=list)
    missing_from_allotted: List[Tuple[str, str]] = field(default_factory=list)

    @property
    def sections(self) -> List[Section]:
        keys = sorted({a.section.key for a in self.assignments})
        return [Section(sem, letter) for sem, letter in keys]

    @property
    def faculty_names(self) -> List[str]:
        return sorted({a.faculty for a in self.assignments})


# ----------------------------------------------------------------------
# sheet parsers
# ----------------------------------------------------------------------
def parse_courses(ws: Worksheet, warnings: List[str]) -> List[Course]:
    """Read the credit catalogue.

    Course codes carry their semester in the first digit of the trailing
    three digit block (``1BAIAK104`` -> 1, ``BCINS701`` -> 7).  Where the
    sheet also carries an explicit ``N SEMESTER`` label, the label is used
    to cross-check the code.
    """

    courses: List[Course] = []
    label_by_row: Dict[int, int] = {}
    current_label: Optional[int] = None

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = clean_text(ws.cell(row=row, column=col).value)
            if not value:
                continue
            sem = parse_semester(value) if re.match(r"^[IVX]+\s*SEMESTER$", value.upper()) else None
            if sem is not None:
                current_label = sem
                label_by_row[row] = sem

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            raw = ws.cell(row=row, column=col).value
            code = clean_text(raw).upper()
            if not code or not COURSE_CODE_RE.match(code):
                continue
            digits = re.search(r"(\d{3})[A-Z]?$", code)
            if not digits:
                continue
            sem_from_code = int(digits.group(1)[0])

            name = clean_text(ws.cell(row=row, column=col + 1).value)
            if not name:
                continue
            credits = ws.cell(row=row, column=col + 2).value
            if not isinstance(credits, (int, float)) or isinstance(credits, bool):
                continue

            label = None
            for label_row, label_sem in label_by_row.items():
                if label_row < row:
                    label = label_sem
            semester = sem_from_code
            if label is not None and label != sem_from_code:
                warnings.append(
                    f"Courses sheet: '{code}' is under a '{semester_ordinal(label)} semester' "
                    f"label but its code implies {semester_ordinal(sem_from_code)}; code wins."
                )
            courses.append(Course(code=code, name=name, credits=float(credits), semester=semester))
    return courses


def _locate_subject_wise_header(ws: Worksheet) -> Tuple[int, int, int, Dict[str, int]]:
    header_row = 0
    section_cols: Dict[str, int] = {}
    for row in range(1, min(ws.max_row, 25) + 1):
        found: Dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            match = SEC_RE.match(clean_text(ws.cell(row=row, column=col).value))
            if match:
                found[match.group(1).upper()] = col
        if len(found) >= 2:
            header_row = row
            section_cols = found
            break

    if not header_row:
        header_row = 1
        section_cols = {letter: 3 + i for i, letter in enumerate("ABCDEF")}

    subject_col = 2
    sem_col = 3
    for col in range(1, ws.max_column + 1):
        head = clean_text(ws.cell(row=header_row, column=col).value).upper()
        if not head:
            continue
        if "SUBJECT" in head or "COURSE" in head:
            subject_col = col
        elif "SEM" in head and "SEC" not in head:
            sem_col = col
    if subject_col in section_cols.values():
        subject_col = 2
    if sem_col in section_cols.values():
        sem_col = 3
    return header_row, subject_col, sem_col, section_cols


def parse_subject_wise(ws: Worksheet, warnings: List[str]) -> List[SubjectRow]:
    """Read the subject -> semester -> section -> faculty grid."""

    header_row, subject_col, sem_col, section_cols = _locate_subject_wise_header(ws)
    rows: List[SubjectRow] = []
    last_semester: Optional[int] = None
    seen: Set[Tuple[str, int, str]] = set()

    for row in range(header_row + 1, ws.max_row + 1):
        raw_sem = clean_text(ws.cell(row=row, column=sem_col).value)
        if raw_sem:
            parsed = parse_semester(raw_sem)
            if parsed is not None:
                last_semester = parsed
        subject = clean_text(ws.cell(row=row, column=subject_col).value)
        if not subject:
            continue
        if last_semester is None:
            warnings.append(
                f"Subject wise row {row}: '{subject}' has no semester label and was skipped."
            )
            continue

        entry = SubjectRow(subject=subject, semester=last_semester)
        for letter, col in section_cols.items():
            faculty = clean_text(ws.cell(row=row, column=col).value)
            if is_blank_marker(faculty):
                continue
            faculty = re.sub(r"\s+", " ", faculty)
            key = (subject, last_semester, letter)
            if key in seen:
                continue
            seen.add(key)
            entry.sections[letter] = faculty
        if entry.sections:
            rows.append(entry)
        else:
            warnings.append(
                f"Subject wise row {row}: '{subject}' has no faculty assigned in any section."
            )
    return rows


def _allotted_groups(ws: Worksheet) -> Tuple[int, List[Tuple[int, int, Optional[int]]]]:
    """Return (header row, [(col_from, col_to, semester_hint)]) from merges."""

    groups: List[Tuple[int, int, Optional[int]]] = []
    header_row = 0
    for rng in ws.merged_cells.ranges:
        if rng.min_row != rng.max_row:
            continue
        label = clean_text(ws.cell(row=rng.min_row, column=rng.min_col).value)
        sem = parse_semester(label) if label else None
        if sem is None:
            continue
        groups.append((rng.min_col, rng.max_col, sem))
        header_row = max(header_row, rng.min_row)
    groups.sort()
    return header_row, groups


def parse_allotted_new(ws: Worksheet, warnings: List[str]) -> Dict[str, FacultyProfile]:
    """Read the faculty load sheet.

    The sheet mixes subject names and numeric workload units inside the same
    semester columns, so numeric cells are treated as workload counters and
    text cells as subjects.
    """

    header_row, groups = _allotted_groups(ws)
    if not groups:
        warnings.append("Allotted New: no semester header merges found, sheet ignored.")
        return {}

    name_col = 2
    for row in (header_row + 1, header_row):
        for col in range(1, ws.max_column + 1):
            if clean_text(ws.cell(row=row, column=col).value).lower() == "name":
                name_col = col
                break

    profiles: Dict[str, FacultyProfile] = {}
    for row in range(header_row + 2, ws.max_row + 1):
        raw_name = clean_text(ws.cell(row=row, column=name_col).value)
        subjects: List[str] = []
        hints: Dict[str, List[int]] = {}
        workload: Optional[float] = None

        for col_from, col_to, sem_hint in groups:
            for col in range(col_from, col_to + 1):
                value = ws.cell(row=row, column=col).value
                if value is None or (isinstance(value, float) and value != value):
                    continue
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    workload = float(value)
                    continue
                subject = clean_text(value)
                if not subject or is_blank_marker(subject):
                    continue
                subjects.append(subject)
                if sem_hint is not None:
                    hints.setdefault(subject, [])
                    if sem_hint not in hints[subject]:
                        hints[subject].append(sem_hint)

        if not subjects and not raw_name:
            continue

        name = raw_name or f"Unnamed Faculty (row {row})"
        if not raw_name:
            warnings.append(
                f"Allotted New row {row}: faculty name is blank; recorded as '{name}'."
            )
        name = re.sub(r"\s+", " ", name)

        profile = profiles.get(name)
        if profile is None:
            profile = FacultyProfile(
                name=name,
                is_placeholder=bool(re.match(r"^new\s+faculty\s*\d*$", name.lower())),
            )
            profiles[name] = profile

        for subject in subjects:
            if subject not in profile.allotted_subjects:
                profile.allotted_subjects.append(subject)
            for sem in hints.get(subject, []):
                bucket = profile.subjects_by_semester.setdefault(sem, [])
                if subject not in bucket:
                    bucket.append(subject)
        if workload is not None:
            profile.workload_units = workload
    return profiles


# ----------------------------------------------------------------------
# hours resolution
# ----------------------------------------------------------------------
class HoursResolver:
    """Map a subject name to its weekly teaching hours."""

    def __init__(self, courses: Sequence[Course], config: TimetableConfig) -> None:
        self.config = config
        self.by_semester: Dict[int, List[Course]] = {}
        self.global_courses: List[Course] = []
        for course in courses:
            self.by_semester.setdefault(course.semester, []).append(course)
            self.global_courses.append(course)
        self.report: Dict[str, Tuple[float, str]] = {}

    def resolve(self, subject: str, semester: int) -> Tuple[float, str]:
        key = subject_key(subject)
        override_key = normalise_name(subject)

        rule_key = (semester, override_key)
        if rule_key in self.config.subject_hour_rules:
            hours = self.config.subject_hour_rules[rule_key]
            self.report[subject] = (hours, "rule")
            return hours, "rule"

        if override_key in self.config.subject_hours_overrides:
            hours = self.config.subject_hours_overrides[override_key]
            self.report[subject] = (hours, "override")
            return hours, "override"

        match = self._match(key, self.by_semester.get(semester, []))
        if match is not None:
            self.report[subject] = (match.credits, f"course:{match.code}")
            return match.credits, f"course:{match.code}"

        match = self._match(key, self.global_courses)
        if match is not None:
            self.report[subject] = (match.credits, f"course:{match.code}")
            return match.credits, f"course:{match.code}"

        if is_lab_subject(subject):
            hours = 2.0
        else:
            hours = 3.0
        self.report[subject] = (hours, "default")
        return hours, "default"

    @staticmethod
    def _match(key: frozenset, pool: Iterable[Course]) -> Optional[Course]:
        best: Optional[Course] = None
        best_score = 0.0
        for course in pool:
            score = jaccard(key, subject_key(course.name))
            if score >= 0.999:
                return course
            if score > best_score:
                best_score = score
                best = course
        if best is not None and best_score >= 0.65:
            return best
        return None


# ----------------------------------------------------------------------
# top level
# ----------------------------------------------------------------------
def parse_workbook(config: TimetableConfig) -> ParseResult:
    """Parse the workbook and build the faculty -> subject -> section map."""

    path = Path(config.excel_path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = load_workbook(path, data_only=True, read_only=False)
    result = ParseResult()

    if config.sheet_courses in workbook.sheetnames:
        result.courses = parse_courses(workbook[config.sheet_courses], result.warnings)
    else:
        result.warnings.append(f"Courses sheet '{config.sheet_courses}' is missing.")

    if config.sheet_subject_wise in workbook.sheetnames:
        result.subject_rows = parse_subject_wise(workbook[config.sheet_subject_wise], result.warnings)
    else:
        raise KeyError(f"Sheet '{config.sheet_subject_wise}' not found in {path.name}")

    if config.sheet_allotted in workbook.sheetnames:
        result.profiles = parse_allotted_new(workbook[config.sheet_allotted], result.warnings)
    else:
        result.warnings.append(f"Allotted sheet '{config.sheet_allotted}' is missing.")

    resolver = HoursResolver(result.courses, config)

    for row in result.subject_rows:
        for letter, faculty in row.sections.items():
            hours, source = resolver.resolve(row.subject, row.semester)
            if hours <= 0:
                result.warnings.append(
                    f"'{row.subject}' resolved to {hours} hours and was skipped."
                )
                continue
            result.assignments.append(
                Assignment(
                    subject=row.subject,
                    section=Section(row.semester, letter),
                    faculty=config.apply_faculty_aliases(faculty),
                    hours=hours,
                    is_lab=is_lab_subject(row.subject),
                    hours_source=source,
                )
            )

    result.hours_report = resolver.report
    _cross_check(result, config)
    workbook.close()
    return result


def _cross_check(result: ParseResult, config: TimetableConfig) -> None:
    """Compare the allotment sheet with the section grid and log the gaps."""

    def same(left: frozenset, right: frozenset) -> bool:
        return left == right or jaccard(left, right) >= 0.8

    allotted: Dict[str, Set[frozenset]] = {}
    for name, profile in result.profiles.items():
        alias = config.apply_faculty_aliases(name)
        allotted[alias] = {subject_key(s) for s in profile.allotted_subjects}

    for name, profile in result.profiles.items():
        alias = config.apply_faculty_aliases(name)
        profile.name = alias
        grid_keys = {subject_key(a.subject) for a in result.assignments if a.faculty == alias}
        for subject in profile.allotted_subjects:
            key = subject_key(subject)
            if not any(same(key, other) for other in grid_keys):
                result.unmapped_subjects.append((alias, subject))

    reported: Set[str] = set()
    for assignment in result.assignments:
        faculty = assignment.faculty
        if faculty not in allotted:
            if faculty not in reported:
                reported.add(faculty)
                result.missing_from_allotted.append((faculty, "<not present in Allotted New>"))
            continue
        key = subject_key(assignment.subject)
        if not any(same(key, other) for other in allotted[faculty]):
            result.missing_from_allotted.append((faculty, assignment.subject))
