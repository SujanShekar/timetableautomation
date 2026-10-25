"""Rendering of the timetable: text report, Excel workbook and CSV slot table."""

from __future__ import annotations

import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import TimetableConfig
from .models import Assignment, Placement, Section, semester_ordinal
from .parser import ParseResult
from .solver import SolveResult
from .validator import ValidationReport

RULE = "=" * 96
THIN = "-" * 96

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
LAB_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


class TimetableView:
    """Convenience index over the final placements."""

    def __init__(self, placements: Sequence[Placement], config: TimetableConfig) -> None:
        self.config = config
        self.placements = list(placements)
        self.by_section: Dict[Tuple[int, str], List[Placement]] = defaultdict(list)
        self.by_faculty: Dict[str, List[Placement]] = defaultdict(list)
        for placement in self.placements:
            meeting = placement.meeting
            self.by_section[meeting.section.key].append(placement)
            self.by_faculty[meeting.faculty].append(placement)
        for bucket in list(self.by_section.values()) + list(self.by_faculty.values()):
            bucket.sort(key=lambda p: (p.day, p.start))

    def section_grid(self, section: Section) -> Dict[int, Dict[int, Placement]]:
        grid: Dict[int, Dict[int, Placement]] = defaultdict(dict)
        for placement in self.by_section.get(section.key, []):
            for period in placement.periods:
                grid[placement.day][period] = placement
        return grid

    def faculty_grid(self, faculty: str) -> Dict[int, Dict[int, Placement]]:
        grid: Dict[int, Dict[int, Placement]] = defaultdict(dict)
        for placement in self.by_faculty.get(faculty, []):
            for period in placement.periods:
                grid[placement.day][period] = placement
        return grid

    @property
    def sections(self) -> List[Section]:
        keys = sorted(self.by_section.keys())
        return [Section(sem, letter) for sem, letter in keys]

    @property
    def faculties(self) -> List[str]:
        return sorted(self.by_faculty.keys())

    @property
    def faculty_names(self) -> List[str]:
        return self.faculties


# ----------------------------------------------------------------------
# text report
# ----------------------------------------------------------------------
def render_text_report(
    config: TimetableConfig,
    parse: ParseResult,
    solve: SolveResult,
    view: TimetableView,
    validation: ValidationReport,
) -> str:
    out: List[str] = []
    add = out.append

    add(RULE)
    add(" ODD SEMESTER 2026 - AUTOMATED TIMETABLE GENERATOR")
    add(RULE)
    add(f" Source workbook : {Path(config.excel_path).name}")
    add(
        f" Sheets read     : {config.sheet_subject_wise}, {config.sheet_allotted}, {config.sheet_courses}"
    )
    add(f" Generated       : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add(
        f" Working week    : {' '.join(config.days)}  "
        f"({config.n_days} days x {config.periods_per_day} periods)"
    )
    add(f" 1st hour limit  : {config.max_first_hour} session(s) per faculty per week")
    add(f" Random seed     : {config.seed}")
    add(
        f" Solver          : {solve.status} in {solve.elapsed_s:.2f}s "
        f"(restarts {solve.attempts}, nodes {solve.nodes}, quality cost {solve.cost:.1f})"
    )
    add("")

    # ---------------- parse summary ----------------
    add(THIN)
    add(" 1. ALLOTMENT SUMMARY")
    add(THIN)
    total_hours = sum(a.total_hours for a in parse.assignments)
    add(f" Subjects in grid            : {len(parse.subject_rows)}")
    add(f" Section-faculty allotments  : {len(parse.assignments)}")
    add(f" Sections scheduled          : {len(view.sections)}")
    add(f" Faculty scheduled           : {len(view.faculties)}")
    add(f" Weekly teaching hours       : {total_hours}")
    add(f" Scheduled class sessions    : {len(view.placements)}")
    add("")

    add(" Section load (hours per week):")
    for section in view.sections:
        hours = sum(
            p.meeting.duration for p in view.by_section.get(section.key, [])
        )
        subjects = len({p.meeting.subject for p in view.by_section.get(section.key, [])})
        add(f"   {section.label:<24} {hours:>3} h   {subjects:>2} subjects")
    add("")

    add(" Faculty load (hours per week):")
    for faculty in view.faculties:
        hours = sum(p.meeting.duration for p in view.by_faculty.get(faculty, []))
        first = validation.faculty_first_hour.get(faculty, 0)
        flag = "  <-- ABOVE LIMIT" if first > config.max_first_hour else ""
        add(
            f"   {faculty:<28} {hours:>3} h   1st hour x{first}{flag}"
        )
    add("")

    # ---------------- section view ----------------
    add(THIN)
    add(" 2. TIMETABLE BY SEMESTER AND SECTION")
    add(THIN)
    for section in view.sections:
        add(f" {section.label}")
        placements = view.by_section.get(section.key, [])
        for day_index, day_name in enumerate(config.days):
            day_placements = sorted(
                (p for p in placements if p.day == day_index), key=lambda p: p.start
            )
            if not day_placements:
                continue
            add(f"   {day_name}")
            for placement in day_placements:
                subject = placement.meeting.subject
                faculty = placement.meeting.faculty
                span = config.period_span(placement.start, placement.end)
                note = "  (lab block)" if placement.meeting.kind == "lab" else ""
                add(f"     {span:<26} {subject:<34} {faculty}{note}")
        add("")

    # ---------------- faculty view ----------------
    add(THIN)
    add(" 3. TIMETABLE BY FACULTY (overlap check)")
    add(THIN)
    for faculty in view.faculties:
        hours = sum(p.meeting.duration for p in view.by_faculty.get(faculty, []))
        first = validation.faculty_first_hour.get(faculty, 0)
        add(f" {faculty}   [{hours} h/week, 1st hour x{first}/{config.max_first_hour}]")
        placements = view.by_faculty.get(faculty, [])
        for day_index, day_name in enumerate(config.days):
            day_placements = sorted(
                (p for p in placements if p.day == day_index), key=lambda p: p.start
            )
            if not day_placements:
                continue
            add(f"   {day_name}")
            for placement in day_placements:
                meeting = placement.meeting
                span = config.period_span(placement.start, placement.end)
                add(
                    f"     {span:<26} {meeting.subject:<34} "
                    f"{meeting.section.label}"
                )
        add("")

    # ---------------- validation ----------------
    add(THIN)
    add(" 4. VALIDATION REPORT")
    add(THIN)
    add(f" Placements checked           : {validation.checked_placements}")
    add(f" Faculty overlap conflicts   : {len([v for v in validation.violations if v.kind == 'faculty overlap'])}")
    add(f" Section overlap conflicts   : {len([v for v in validation.violations if v.kind == 'section overlap'])}")
    add(f" 1st hour cap violations     : {len([v for v in validation.violations if v.kind == 'first hour cap'])}")
    add(f" Coverage errors             : {len([v for v in validation.violations if v.kind == 'coverage'])}")
    add(f" RESULT                      : {'PASS - timetable is conflict free' if validation.ok else 'FAIL'}")
    add("")
    if validation.violations:
        add(" Violations:")
        for violation in validation.violations:
            add(f"   - {violation}")
        add("")

    add(" 1st hour distribution (faculty -> number of 1st hour sessions):")
    histogram: Dict[int, List[str]] = defaultdict(list)
    for faculty, count in validation.faculty_first_hour.items():
        histogram[count].append(faculty)
    for count in sorted(histogram):
        add(f"   x{count}: {len(histogram[count])} faculty -> {', '.join(sorted(histogram[count]))}")
    add("")

    # ---------------- master mapping ----------------
    add(THIN)
    add(" 5. FACULTY -> SUBJECT -> SECTION MASTER")
    add(THIN)
    for faculty in view.faculties:
        add(f" {faculty}")
        grouped: Dict[Tuple[int, str, str], List[Placement]] = defaultdict(list)
        for placement in view.by_faculty.get(faculty, []):
            key = (
                placement.meeting.section.semester,
                placement.meeting.section.letter,
                placement.meeting.subject,
            )
            grouped[key].append(placement)
        for key in sorted(grouped):
            semester, letter, subject = key
            placements = sorted(grouped[key], key=lambda p: (p.day, p.start))
            hours = sum(p.meeting.duration for p in placements)
            slots = ", ".join(
                f"{config.days[p.day]} P{p.start + 1}"
                + (f"-P{p.end + 1}" if p.meeting.duration > 1 else "")
                for p in placements
            )
            label = f"{semester_ordinal(semester)} Sem - Sec {letter}"
            add(f"     {label:<24} {subject:<36} {hours:>2}h   {slots}")
        add("")

    # ---------------- data notes ----------------
    add(THIN)
    add(" 6. DATA NOTES AND WARNINGS")
    add(THIN)
    if not parse.warnings and not parse.unmapped_subjects and not parse.missing_from_allotted:
        add(" No data quality issues detected.")
    for warning in parse.warnings:
        add(f" ! {warning}")

    unmapped: Dict[str, List[str]] = defaultdict(list)
    for faculty, subject in parse.unmapped_subjects:
        unmapped[faculty].append(subject)
    for faculty in sorted(unmapped):
        add(
            f" ! Allotted but not mapped to any section: {faculty} -> "
            + ", ".join(sorted(set(unmapped[faculty])))
        )

    missing: Dict[str, List[str]] = defaultdict(list)
    for faculty, subject in parse.missing_from_allotted:
        missing[faculty].append(subject)
    for faculty in sorted(missing):
        add(
            f" ! Mapped in 'Subject wise' but absent from 'Allotted New': {faculty} -> "
            + ", ".join(sorted(set(missing[faculty])))
        )
    default_hours = sorted(
        subject for subject, (_, source) in parse.hours_report.items() if source == "default"
    )
    if default_hours:
        add(" ! Subjects using the default weekly hours (edit config or use --hours-file):")
        for subject in default_hours:
            add(f"     - {subject}: {parse.hours_report[subject][0]:g} h/week")
    placeholders = sorted(
        {
            a.faculty
            for a in parse.assignments
            if a.faculty.lower().startswith("new faculty")
        }
    )
    if placeholders:
        add(
            " ! Placeholder faculty names still present: "
            + ", ".join(placeholders)
            + " (map them with --alias-file to real names)"
        )
    add("")
    add(RULE)
    add(" End of report")
    add(RULE)
    return "\n".join(out)


# ----------------------------------------------------------------------
# excel report
# ----------------------------------------------------------------------
def _style_header_row(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def write_excel_report(
    path: Path,
    config: TimetableConfig,
    parse: ParseResult,
    solve: SolveResult,
    view: TimetableView,
    validation: ValidationReport,
) -> Path:
    workbook = Workbook()

    # ---- section timetable ----
    ws = workbook.active
    ws.title = "Timetable by Section"
    row = 1
    for semester in sorted({s.semester for s in view.sections}):
        ws.cell(row=row, column=1, value=f"{semester_ordinal(semester)} SEMESTER").font = TITLE_FONT
        row += 1
        for section in [s for s in view.sections if s.semester == semester]:
            ws.cell(row=row, column=1, value=section.label).font = BOLD
            ws.cell(row=row, column=1).fill = SUB_FILL
            for col in range(2, 2 + config.n_days):
                ws.cell(row=row, column=col).fill = SUB_FILL
            row += 1
            header = ["Day"] + [f"P{i + 1}\n{config.period_label(i)}" for i in range(config.periods_per_day)]
            for col, text in enumerate(header, start=1):
                ws.cell(row=row, column=col, value=text)
            _style_header_row(ws, row, len(header))
            row += 1
            grid = view.section_grid(section)
            for day_index, day_name in enumerate(config.days):
                ws.cell(row=row, column=1, value=day_name).font = BOLD
                slots = grid.get(day_index, {})
                for period in range(config.periods_per_day):
                    placement = slots.get(period)
                    cell = ws.cell(row=row, column=2 + period)
                    if placement is not None:
                        meeting = placement.meeting
                        cell.value = f"{meeting.subject}\n({meeting.faculty})"
                        if meeting.kind == "lab":
                            cell.fill = LAB_FILL
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.border = THIN_BORDER
                row += 1
            row += 1
    ws.column_dimensions["A"].width = 18
    for col in range(2, 2 + config.n_days):
        ws.column_dimensions[get_column_letter(col)].width = 26
    ws.freeze_panes = "B2"

    # ---- faculty timetable ----
    ws = workbook.create_sheet("Timetable by Faculty")
    row = 1
    for faculty in view.faculties:
        hours = sum(p.meeting.duration for p in view.by_faculty.get(faculty, []))
        first = validation.faculty_first_hour.get(faculty, 0)
        ws.cell(row=row, column=1, value=f"{faculty}   ({hours} h/week, 1st hour x{first})").font = TITLE_FONT
        row += 1
        header = ["Day"] + [f"P{i + 1}\n{config.period_label(i)}" for i in range(config.periods_per_day)]
        for col, text in enumerate(header, start=1):
            ws.cell(row=row, column=col, value=text)
        _style_header_row(ws, row, len(header))
        row += 1
        grid = view.faculty_grid(faculty)
        for day_index, day_name in enumerate(config.days):
            ws.cell(row=row, column=1, value=day_name).font = BOLD
            slots = grid.get(day_index, {})
            for period in range(config.periods_per_day):
                placement = slots.get(period)
                cell = ws.cell(row=row, column=2 + period)
                if placement is not None:
                    meeting = placement.meeting
                    cell.value = f"{meeting.subject}\n{meeting.section.label}"
                    if meeting.kind == "lab":
                        cell.fill = LAB_FILL
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = THIN_BORDER
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 26
    for col in range(2, 2 + config.n_days):
        ws.column_dimensions[get_column_letter(col)].width = 26
    ws.freeze_panes = "B2"

    # ---- section summary ----
    ws = workbook.create_sheet("Section Summary")
    header = ["Semester", "Section", "Subjects", "Hours / week", "Days used", "First hour", "1st-hour holder"]
    for col, text in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=text)
    _style_header_row(ws, 1, len(header))
    row = 2
    for section in view.sections:
        placements = view.by_section.get(section.key, [])
        hours = sum(p.meeting.duration for p in placements)
        subjects = sorted({p.meeting.subject for p in placements})
        first_holder = next(
            (
                f"{p.meeting.faculty} ({config.days[p.day]})"
                for p in placements
                if p.start == 0
            ),
            "-",
        )
        ws.cell(row=row, column=1, value=section.semester)
        ws.cell(row=row, column=2, value=section.label)
        ws.cell(row=row, column=3, value=len(subjects))
        ws.cell(row=row, column=4, value=hours)
        ws.cell(row=row, column=5, value=validation.section_days.get(section.key, 0))
        ws.cell(row=row, column=6, value=sum(1 for p in placements if p.start == 0))
        ws.cell(row=row, column=7, value=first_holder)
        row += 1
    for col, width in enumerate([12, 24, 12, 14, 12, 12, 32], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # ---- faculty load ----
    ws = workbook.create_sheet("Faculty Load")
    header = [
        "Faculty",
        "Semesters",
        "Subjects",
        "Hours / week",
        "Days used",
        "1st hour sessions",
        "Limit",
        "Sections",
    ]
    for col, text in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=text)
    _style_header_row(ws, 1, len(header))
    row = 2
    for faculty in view.faculties:
        placements = view.by_faculty.get(faculty, [])
        hours = sum(p.meeting.duration for p in placements)
        subjects = sorted({p.meeting.subject for p in placements})
        semesters = sorted({p.meeting.section.semester for p in placements})
        sections = sorted({p.meeting.section.label for p in placements})
        ws.cell(row=row, column=1, value=faculty)
        ws.cell(row=row, column=2, value=", ".join(semester_ordinal(s) for s in semesters))
        ws.cell(row=row, column=3, value="; ".join(subjects))
        ws.cell(row=row, column=4, value=hours)
        ws.cell(row=row, column=5, value=validation.faculty_days.get(faculty, 0))
        ws.cell(row=row, column=6, value=validation.faculty_first_hour.get(faculty, 0))
        ws.cell(row=row, column=7, value=config.max_first_hour)
        ws.cell(row=row, column=8, value="; ".join(sections))
        row += 1
    for col, width in enumerate([26, 16, 44, 14, 12, 16, 8, 40], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # ---- subject master ----
    ws = workbook.create_sheet("Subject Master")
    header = ["Subject", "Semester", "Section", "Faculty", "Hours / week", "Type", "Course code", "Hours from"]
    for col, text in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=text)
    _style_header_row(ws, 1, len(header))
    row = 2
    for assignment in sorted(
        parse.assignments,
        key=lambda a: (a.section.semester, a.section.letter, a.subject),
    ):
        ws.cell(row=row, column=1, value=assignment.subject)
        ws.cell(row=row, column=2, value=assignment.section.label)
        ws.cell(row=row, column=3, value=assignment.section.letter)
        ws.cell(row=row, column=4, value=assignment.faculty)
        ws.cell(row=row, column=5, value=assignment.total_hours)
        ws.cell(row=row, column=6, value="Lab" if assignment.is_lab else "Theory")
        ws.cell(row=row, column=7, value=assignment.course_code or "-")
        ws.cell(row=row, column=8, value=assignment.hours_source)
        row += 1
    for col, width in enumerate([38, 22, 10, 28, 14, 10, 16, 22], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # ---- validation ----
    ws = workbook.create_sheet("Validation")
    ws.cell(row=1, column=1, value="Check").font = BOLD
    ws.cell(row=1, column=2, value="Result").font = BOLD
    ws.cell(row=1, column=3, value="Detail").font = BOLD
    _style_header_row(ws, 1, 3)
    checks = [
        (
            "Faculty overlap",
            "PASS" if not [v for v in validation.violations if v.kind == "faculty overlap"] else "FAIL",
            f"{len([v for v in validation.violations if v.kind == 'faculty overlap'])} conflict(s)",
        ),
        (
            "Section overlap",
            "PASS" if not [v for v in validation.violations if v.kind == "section overlap"] else "FAIL",
            f"{len([v for v in validation.violations if v.kind == 'section overlap'])} conflict(s)",
        ),
        (
            "1st hour distribution",
            "PASS"
            if not [v for v in validation.violations if v.kind == "first hour cap"]
            else "FAIL",
            f"limit {config.max_first_hour} per faculty, "
            f"{len(validation.faculty_first_hour)} faculty hold a 1st hour",
        ),
        (
            "Hour coverage",
            "PASS" if not [v for v in validation.violations if v.kind == "coverage"] else "FAIL",
            f"{len([v for v in validation.violations if v.kind == 'coverage'])} unscheduled allotment(s)",
        ),
        (
            "Grid bounds",
            "PASS" if not [v for v in validation.violations if v.kind == "grid"] else "FAIL",
            f"{config.n_days} days x {config.periods_per_day} periods",
        ),
    ]
    row = 2
    for name, result, detail in checks:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=result)
        ws.cell(row=row, column=3, value=detail)
        if result == "FAIL":
            ws.cell(row=row, column=2).font = Font(bold=True, color="C00000")
        row += 1
    if validation.violations:
        row += 1
        ws.cell(row=row, column=1, value="Violations").font = BOLD
        row += 1
        for violation in validation.violations:
            ws.cell(row=row, column=1, value=violation.kind)
            ws.cell(row=row, column=2, value=violation.severity)
            ws.cell(row=row, column=3, value=violation.detail)
            row += 1
    for col, width in enumerate([26, 12, 90], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # ---- data notes ----
    ws = workbook.create_sheet("Data Notes")
    ws.cell(row=1, column=1, value="Type").font = BOLD
    ws.cell(row=1, column=2, value="Detail").font = BOLD
    _style_header_row(ws, 1, 2)
    row = 2
    notes: List[Tuple[str, str]] = [("warning", w) for w in parse.warnings]
    notes += [("allotment gap", f"{f} -> {s}") for f, s in parse.unmapped_subjects]
    notes += [("subject-wise gap", f"{f} -> {s}") for f, s in parse.missing_from_allotted]
    notes += [
        ("default hours", f"{s}: {parse.hours_report[s][0]:g} h/week")
        for s in sorted(parse.hours_report)
        if parse.hours_report[s][1] == "default"
    ]
    if not notes:
        notes.append(("info", "no data quality issues detected"))
    for kind, detail in notes:
        ws.cell(row=row, column=1, value=kind)
        ws.cell(row=row, column=2, value=detail)
        row += 1
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


# ----------------------------------------------------------------------
# csv slot table
# ----------------------------------------------------------------------
def write_csv(path: Path, config: TimetableConfig, view: TimetableView) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["semester", "section", "day", "period", "time", "subject", "faculty", "type"]
        )
        for placement in sorted(view.placements, key=lambda p: (p.meeting.section.key, p.day, p.start)):
            meeting = placement.meeting
            writer.writerow(
                [
                    f"{semester_ordinal(meeting.section.semester)} Sem",
                    f"Sec {meeting.section.letter}",
                    config.days[placement.day],
                    placement.start + 1,
                    config.period_label(placement.start),
                    meeting.subject,
                    meeting.faculty,
                    "Lab" if meeting.kind == "lab" else "Theory",
                ]
            )
    return path
